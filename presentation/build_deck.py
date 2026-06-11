"""insta_trend 시계열 분석 PPT 빌더 (실데이터판).

실행::

    python presentation/build_deck.py

네이버 데이터랩 실데이터(``naver_datalab_panel.csv``)로 차트를 다시 그리고
``insta_trend_timeseries.pptx`` 를 조립합니다.

구성: CS 시계열 수업 이론(정상성·ACF·분해·ARIMA) → 실데이터 분석 →
2026-06-11 기준 한국 트렌딩 해시태그 결론 + 활용 전략 + 핵심 코드 설명.
폰트는 임베드하지 않으며 한글 본문은 맑은 고딕을 요청합니다.
"""

from __future__ import annotations

from itertools import count
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from charts import FOCUS, render_all

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
OUT = HERE / "insta_trend_timeseries.pptx"

KO_FONT = "맑은 고딕"
EN_FONT = "Segoe UI"
CODE_FONT = "Consolas"

PINK = RGBColor(0xE1, 0x30, 0x6C)
PURPLE = RGBColor(0x83, 0x3A, 0xB4)
INK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
LIGHT = RGBColor(0xF2, 0xF2, 0xF4)
CODE_BG = RGBColor(0x2B, 0x2B, 0x33)
CODE_FG = RGBColor(0xE8, 0xE8, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SW, SH = Inches(13.333), Inches(7.5)

ASOF = "2026년 6월 11일 (KST)"


def _set_font(run, size, *, bold=False, color=INK, font=KO_FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _box(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf


def _rect(slide, left, top, width, height, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _footer(slide, page):
    _rect(slide, 0, SH - Inches(0.32), SW, Inches(0.32), LIGHT)
    _, tf = _box(slide, Inches(0.4), SH - Inches(0.34), Inches(8), Inches(0.3))
    r = tf.paragraphs[0].add_run()
    r.text = "시계열 데이터 분석 · 한국 해시태그 트렌드 · 기준일 2026-06-11"
    _set_font(r, 9, color=GREY)
    _, tf2 = _box(slide, SW - Inches(1.2), SH - Inches(0.34), Inches(0.9), Inches(0.3))
    p = tf2.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r2 = p.add_run()
    r2.text = str(page)
    _set_font(r2, 9, color=GREY)


def header(slide, title, eyebrow=None):
    _rect(slide, Inches(0.4), Inches(0.42), Inches(0.18), Inches(0.55), PINK)
    if eyebrow:
        _, tfe = _box(slide, Inches(0.75), Inches(0.28), Inches(11), Inches(0.35))
        re_ = tfe.paragraphs[0].add_run()
        re_.text = eyebrow.upper()
        _set_font(re_, 11, bold=True, color=PURPLE, font=EN_FONT)
    _, tf = _box(slide, Inches(0.72), Inches(0.5), Inches(11.5), Inches(0.7))
    r = tf.paragraphs[0].add_run()
    r.text = title
    _set_font(r, 28, bold=True, color=INK)


# --------------------------------------------------------------------------- #
# 슬라이드 타입
# --------------------------------------------------------------------------- #
def title_slide(prs, stats):
    s = _blank(prs)
    _rect(s, 0, 0, SW, SH, WHITE)
    _rect(s, 0, 0, Inches(0.35), SH, PINK)
    _rect(s, Inches(0.35), 0, Inches(0.12), SH, PURPLE)

    _, tf = _box(s, Inches(1.0), Inches(1.9), Inches(11.5), Inches(2.2))
    r = tf.paragraphs[0].add_run()
    r.text = "한국 인스타그램 해시태그"
    _set_font(r, 30, color=GREY)
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "시계열 데이터 분석"
    _set_font(r2, 54, bold=True, color=INK)

    _, tf2 = _box(s, Inches(1.0), Inches(4.3), Inches(11.3), Inches(1.2))
    r3 = tf2.paragraphs[0].add_run()
    r3.text = "네이버 데이터랩 실데이터로 본 2026년 상반기 해시태그 트렌드"
    _set_font(r3, 18, color=GREY)
    p4 = tf2.add_paragraph()
    r4 = p4.add_run()
    r4.text = "— 정상성 · 자기상관 · 분해 · 예측: 수업에서 배운 시계열 이론의 실전 적용"
    _set_font(r4, 15, color=GREY)

    _, tf3 = _box(s, Inches(1.0), Inches(6.1), Inches(11.3), Inches(0.9))
    r5 = tf3.paragraphs[0].add_run()
    r5.text = (
        f"기준일 {ASOF}  ·  분석 기간 {stats['date_start']} ~ {stats['date_end']}"
        f" ({stats['n_days']}일)  ·  해시태그 {stats['n_tags']}개"
    )
    _set_font(r5, 13, color=PINK, bold=True)
    return s


def agenda_slide(prs, page):
    s = _blank(prs)
    header(s, "목차", "Agenda")
    items = [
        ("01", "문제 정의와 데이터", "마케터의 고민: 무엇이 떡상할까? · 네이버 데이터랩 실데이터"),
        ("02", "이론: 시계열의 구조", "확률과정, 구성요소, 가법/승법 분해 모형"),
        ("03", "이론: 정상성과 차분", "약정상성 3조건, 단위근, 1차 차분"),
        ("04", "이론: 자기상관과 ARIMA", "ACF/PACF, 백색잡음, AR·MA·ARIMA(p,d,q)"),
        ("05", "분석: 평활화 · 정상성 · ACF", "이론을 실데이터로 검증"),
        ("06", "분석: 분해 · 계절성 · 모멘텀 · 예측", "추세/요일 효과 분리와 14일 외삽"),
        ("07", "결론: 지금 뜨는 해시태그", "2026-06-11 기준 트렌딩 태그와 활용 전략"),
        ("08", "부록: 핵심 코드 · 한계 · 재현", "pandas/numpy 구현과 데이터 캐비앳"),
    ]
    top = Inches(1.7)
    for i, (num, t, sub) in enumerate(items):
        row = top + Inches(0.62) * i
        _, tf = _box(s, Inches(0.9), row, Inches(0.9), Inches(0.55))
        r = tf.paragraphs[0].add_run()
        r.text = num
        _set_font(r, 18, bold=True, color=PINK)
        _, tf2 = _box(s, Inches(1.8), row, Inches(5.6), Inches(0.55))
        r2 = tf2.paragraphs[0].add_run()
        r2.text = t
        _set_font(r2, 16, bold=True, color=INK)
        _, tf3 = _box(s, Inches(7.5), row, Inches(5.3), Inches(0.55))
        r3 = tf3.paragraphs[0].add_run()
        r3.text = sub
        _set_font(r3, 12.5, color=GREY)
    _footer(s, page)
    return s


def bullets_slide(prs, page, title, eyebrow, bullets, lead=None):
    s = _blank(prs)
    header(s, title, eyebrow)
    top = Inches(1.5)
    if lead:
        _rect(s, Inches(0.75), top, Inches(11.8), Inches(0.9), LIGHT)
        _, tf = _box(s, Inches(1.0), top + Inches(0.12), Inches(11.3), Inches(0.7))
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = tf.paragraphs[0].add_run()
        r.text = lead
        _set_font(r, 15, bold=True, color=INK)
        top = top + Inches(1.25)
    _, tf = _box(s, Inches(0.9), top, Inches(11.6), Inches(4.6))
    for i, (head, body) in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        r = p.add_run()
        r.text = f"●  {head}"
        _set_font(r, 16.5, bold=True, color=PINK)
        if body:
            p2 = tf.add_paragraph()
            p2.space_after = Pt(13)
            r2 = p2.add_run()
            r2.text = f"     {body}"
            _set_font(r2, 13.5, color=GREY)
    _footer(s, page)
    return s


def theory_slide(prs, page, title, eyebrow, concepts, formulas):
    """이론 슬라이드: 왼쪽 개념 불릿 + 오른쪽 수식 패널."""

    s = _blank(prs)
    header(s, title, eyebrow)

    _, tf = _box(s, Inches(0.8), Inches(1.55), Inches(6.9), Inches(5.3))
    for i, (head, body) in enumerate(concepts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = f"●  {head}"
        _set_font(r, 15.5, bold=True, color=PINK)
        p2 = tf.add_paragraph()
        p2.space_after = Pt(12)
        r2 = p2.add_run()
        r2.text = f"     {body}"
        _set_font(r2, 12.5, color=GREY)

    # 수식 패널
    px, py, pw, ph = Inches(7.95), Inches(1.55), Inches(4.85), Inches(5.3)
    _rect(s, px, py, pw, ph, LIGHT)
    _rect(s, px, py, pw, Inches(0.5), PURPLE)
    _, tfh = _box(s, px + Inches(0.2), py + Inches(0.06), pw - Inches(0.4), Inches(0.4))
    rh = tfh.paragraphs[0].add_run()
    rh.text = "핵심 수식"
    _set_font(rh, 13, bold=True, color=WHITE)
    _, tff = _box(s, px + Inches(0.25), py + Inches(0.7), pw - Inches(0.5), ph - Inches(0.9))
    for i, (label, formula) in enumerate(formulas):
        p = tff.paragraphs[0] if i == 0 else tff.add_paragraph()
        p.space_after = Pt(2)
        r = p.add_run()
        r.text = label
        _set_font(r, 11.5, bold=True, color=PURPLE)
        p2 = tff.add_paragraph()
        p2.space_after = Pt(12)
        r2 = p2.add_run()
        r2.text = formula
        _set_font(r2, 13.5, color=INK, font=CODE_FONT)
    _footer(s, page)
    return s


def chart_slide(prs, page, title, eyebrow, image, notes):
    s = _blank(prs)
    header(s, title, eyebrow)
    s.shapes.add_picture(str(ASSETS / image), Inches(0.55), Inches(1.45), width=Inches(8.5))
    panel_left = Inches(9.25)
    _rect(s, panel_left, Inches(1.45), Inches(3.55), Inches(5.4), LIGHT)
    _, tf = _box(s, panel_left + Inches(0.22), Inches(1.65), Inches(3.15), Inches(5.1))
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = "읽는 법"
    _set_font(r0, 13, bold=True, color=PURPLE)
    p0.space_after = Pt(8)
    for note in notes:
        p = tf.add_paragraph()
        p.space_after = Pt(10)
        r = p.add_run()
        r.text = f"• {note}"
        _set_font(r, 12.5, color=INK)
    _footer(s, page)
    return s


def section_slide(prs, kicker, title):
    s = _blank(prs)
    _rect(s, 0, 0, SW, SH, INK)
    _rect(s, Inches(0.9), Inches(3.0), Inches(1.4), Inches(0.12), PINK)
    _, tf = _box(s, Inches(0.9), Inches(2.2), Inches(11), Inches(0.6))
    r = tf.paragraphs[0].add_run()
    r.text = kicker.upper()
    _set_font(r, 14, bold=True, color=PINK, font=EN_FONT)
    _, tf2 = _box(s, Inches(0.9), Inches(3.3), Inches(11.5), Inches(1.5))
    r2 = tf2.paragraphs[0].add_run()
    r2.text = title
    _set_font(r2, 40, bold=True, color=WHITE)
    return s


def cards_slide(prs, page, title, eyebrow, cards):
    s = _blank(prs)
    header(s, title, eyebrow)
    cw, ch = Inches(5.9), Inches(2.1)
    gap = Inches(0.3)
    lefts = [Inches(0.75), Inches(0.75) + cw + gap]
    tops = [Inches(1.55), Inches(1.55) + ch + gap]
    for i, (tag, big, body, color) in enumerate(cards[:4]):
        left = lefts[i % 2]
        top = tops[i // 2]
        _rect(s, left, top, cw, ch, LIGHT)
        _rect(s, left, top, Inches(0.14), ch, color)
        _, tf = _box(s, left + Inches(0.35), top + Inches(0.18), cw - Inches(0.6), ch - Inches(0.36))
        r = tf.paragraphs[0].add_run()
        r.text = tag
        _set_font(r, 13, bold=True, color=color)
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = big
        _set_font(r2, 22, bold=True, color=INK)
        p3 = tf.add_paragraph()
        p3.space_before = Pt(4)
        r3 = p3.add_run()
        r3.text = body
        _set_font(r3, 12.5, color=GREY)
    _footer(s, page)
    return s


def code_slide(prs, page, title, eyebrow, blocks):
    """핵심 코드 슬라이드: (설명, 코드) 블록 목록."""

    s = _blank(prs)
    header(s, title, eyebrow)
    top = Inches(1.5)
    for desc, code in blocks:
        _, tfd = _box(s, Inches(0.8), top, Inches(11.8), Inches(0.35))
        rd = tfd.paragraphs[0].add_run()
        rd.text = f"●  {desc}"
        _set_font(rd, 13.5, bold=True, color=PINK)
        top += Inches(0.42)
        lines = code.strip("\n").split("\n")
        bh = Inches(0.28) * len(lines) + Inches(0.22)
        _rect(s, Inches(1.0), top, Inches(11.4), bh, CODE_BG)
        _, tfc = _box(s, Inches(1.25), top + Inches(0.1), Inches(11.0), bh - Inches(0.16))
        for j, line in enumerate(lines):
            p = tfc.paragraphs[0] if j == 0 else tfc.add_paragraph()
            r = p.add_run()
            r.text = line
            _set_font(r, 12, color=CODE_FG, font=CODE_FONT)
        top += bh + Inches(0.18)
    _footer(s, page)
    return s


# --------------------------------------------------------------------------- #
def build():
    stats = render_all()
    pg = count(1)
    g = stats["growth_ytd"]
    m = stats["mom_30"]

    prs = Presentation()
    prs.slide_width = Emu(int(SW))
    prs.slide_height = Emu(int(SH))

    title_slide(prs, stats)
    agenda_slide(prs, next(pg))

    bullets_slide(
        prs, next(pg), "문제 정의 — 마케터의 고민", "01 · Problem",
        lead="“다음 달에 어떤 해시태그가 떡상할까? 우리 계정 도달은 오를까, 꺾일까?” — 감으로는 답할 수 없는 질문들",
        bullets=[
            ("“어떤 태그가 떡상할까?”", "인기 순위는 '이미 큰 것'만 보여줌 — 순위에 올라온 시점은 대개 피크 직전이라 따라 들어가면 늦음"),
            ("“우리 계정은 어떻게 될까?”", "도달·팔로워 그래프가 오르는 중인지, 꺾이는 중인지, 그냥 요일 출렁임인지 눈으로는 구분 불가"),
            ("“왜 예측이 어렵나?”", "관심도 = 추세 + 계절성 + 이벤트 + 노이즈가 섞인 신호 — 분리하지 않으면 노이즈를 트렌드로 오독"),
            ("이 발표의 제안", "시계열 분해로 신호를 성분별로 나누면 '진짜 상승'·'피크 통과'·'단순 요일 효과'를 데이터로 구분할 수 있다"),
        ],
    )

    bullets_slide(
        prs, next(pg), "분석 접근", "01 · Approach",
        lead="질문: 2026년 6월 11일 현재, 한국에서 어떤 해시태그가 뜨고 있고 — 어디로 가는가?",
        bullets=[
            ("왜 시계열인가", "한 시점의 인기 순위는 '지금 큰 것'만 보여줌 — 추세·계절성·모멘텀을 분리해야 '뜨는 것'과 '식는 것'을 구분"),
            ("데이터", "네이버 데이터랩 검색어트렌드 실데이터 — 한국 사용자의 일별 검색 관심도 (2026-06-11 수집)"),
            ("프록시 가정", "인스타그램은 해시태그 시계열 공개 API가 없음 → 같은 주제의 네이버 검색량을 한국 관심도의 대리지표로 사용"),
            ("분석 절차", "EDA → 평활화 → 정상성 진단 → ACF → 분해 → 계절성/모멘텀 → 예측 → 액션"),
        ],
    )

    bullets_slide(
        prs, next(pg), "데이터셋", "01 · Dataset",
        lead=f"{stats['date_start']} ~ {stats['date_end']} · {stats['n_days']}일 · 일 단위 · 라이프스타일 해시태그 {stats['n_tags']}개",
        bullets=[
            ("수집", "네이버 데이터랩 검색어트렌드 API — 키워드그룹별 일별 ratio (그룹 예: 러닝=러닝+러닝크루)"),
            ("대상 태그", "#러닝 #캠핑 #등산 #카페 #국내여행 #다이어트 #피크닉 #클라이밍 #빵지순례"),
            ("ratio의 의미", "요청 내 최대값=100인 상대값 → 시리즈 간 절대 비교는 불가, 시리즈 내 시간 변화 분석은 유효"),
            ("형태", "tidy 패널 (date, hashtag, ratio) → pivot하여 날짜×태그 행렬로 분석"),
        ],
    )

    # ---------------- Part 1: 이론 ---------------- #
    section_slide(prs, "Part 1 · Theory", "시계열 이론 — 수업에서 배운 도구들")

    theory_slide(
        prs, next(pg), "시계열 데이터란", "02 · Stochastic process",
        concepts=[
            ("정의", "시간 순서로 관측된 확률변수의 열 — 확률과정 {X_t}의 한 실현(realization)"),
            ("핵심 성질: 시간 의존성", "관측치가 i.i.d.가 아님 — 오늘 값이 어제 값과 상관 → 고전 통계의 독립 가정이 무너짐"),
            ("왜 따로 배우는가", "회귀처럼 다루면 표준오차·검정이 모두 왜곡 — 자기상관을 모형 안에 넣어야 함"),
            ("우리 데이터", "해시태그별 일별 관심도 161개 관측치 = 일 단위 이산 시계열 9개"),
        ],
        formulas=[
            ("확률과정", "{ X_t : t = 1, 2, ..., n }"),
            ("자기공분산", "γ(k) = Cov(X_t, X_{t+k})"),
            ("시간 의존성", "Corr(X_t, X_{t-1}) ≠ 0"),
            ("표본 (우리 데이터)", "n = 161일, 9개 시리즈"),
        ],
    )

    theory_slide(
        prs, next(pg), "구성요소와 분해 모형", "02 · Decomposition",
        concepts=[
            ("4대 구성요소", "추세(T) 장기 방향 · 계절(S) 고정 주기 반복 · 순환(C) 비고정 장주기 · 불규칙(R) 잔차"),
            ("가법 모형", "변동 폭이 수준과 무관하게 일정할 때 — 성분을 더해서 관측을 설명"),
            ("승법 모형", "수준이 커질수록 변동도 커질 때 — 로그 변환하면 가법으로 환원"),
            ("추정 방법", "추세는 중심 이동평균, 계절성분은 '같은 요일끼리 평균', 잔차는 나머지"),
        ],
        formulas=[
            ("가법 모형", "X_t = T_t + S_t + R_t"),
            ("승법 모형", "X_t = T_t × S_t × R_t"),
            ("로그 변환", "log X_t = log T_t + log S_t + log R_t"),
            ("추세 추정 (period=7)", "T_t = (1/7) Σ X_{t+j}, j=-3..3"),
            ("계절 추정", "S_t = mean( X−T | 같은 요일 )"),
        ],
    )

    theory_slide(
        prs, next(pg), "정상성과 차분", "03 · Stationarity",
        concepts=[
            ("약정상성 3조건", "① 평균이 시간에 불변 ② 분산이 유한·불변 ③ 자기공분산이 시차 k에만 의존"),
            ("왜 중요한가", "ARMA류 모형·ACF 해석·예측 이론이 모두 정상성 위에서 성립 — 비정상 계열은 가짜 상관(spurious) 위험"),
            ("진단", "시각: 롤링 평균/분산이 움직이는가 · ACF가 천천히 감쇠하는가 / 검정: ADF 단위근 검정"),
            ("처방: 차분", "1차 차분으로 추세 제거, 계절차분(lag 7)으로 주기 제거 — ARIMA의 d가 바로 차분 횟수"),
        ],
        formulas=[
            ("약정상성", "E[X_t] = μ,  Var(X_t) = σ² < ∞"),
            ("", "Cov(X_t, X_{t+k}) = γ(k)  ∀t"),
            ("1차 차분", "∇X_t = X_t − X_{t-1}"),
            ("계절 차분 (주간)", "∇₇X_t = X_t − X_{t-7}"),
            ("랜덤워크(비정상)", "X_t = X_{t-1} + ε_t"),
        ],
    )

    theory_slide(
        prs, next(pg), "자기상관 · 백색잡음 · ARIMA", "04 · ACF & ARIMA",
        concepts=[
            ("ACF", "시차 k에서 자기 자신과의 상관 ρ(k) — 추세면 느리게 감쇠, 주기 s면 lag s마다 스파이크"),
            ("백색잡음", "평균 0, 상관 0인 순수 잡음 — 잔차가 백색잡음이면 모형이 구조를 다 흡수했다는 신호"),
            ("판정 기준", "표본 ACF가 ±1.96/√n 밴드 안이면 해당 시차 상관은 유의하지 않음 (95%)"),
            ("ARIMA(p,d,q)", "AR(p) 과거 값 회귀 + I(d) 차분 + MA(q) 과거 충격 회귀 — 차수는 ACF/PACF로 식별"),
            ("우리의 선택", "수업 범위의 해석 가능한 기본형: 선형추세 + 요일 더미 → ARIMA 대비 단순하지만 구조가 투명"),
        ],
        formulas=[
            ("자기상관함수", "ρ(k) = γ(k) / γ(0)"),
            ("표본 ACF 한계", "±1.96 / √n  (n=161 → ±0.155)"),
            ("백색잡음", "ε_t ~ WN(0, σ²)"),
            ("AR(p)", "X_t = Σ φ_i X_{t-i} + ε_t"),
            ("MA(q)", "X_t = ε_t + Σ θ_j ε_{t-j}"),
            ("ARIMA", "∇^d X_t 가 ARMA(p,q)"),
        ],
    )

    # ---------------- Part 2: 분석 ---------------- #
    section_slide(prs, "Part 2 · Analysis", "실데이터 분석 — 이론을 데이터에 적용")

    chart_slide(
        prs, next(pg), "원시 시계열 관찰", "05 · Raw signal", "01_overview.png",
        ["9개 태그의 일별 관심도 원계열", "톱니 모양 = 주 7일 계절성", "#캠핑·#피크닉: 봄에 큰 산(계절 추세)", "#카페: 1월 고점 후 하락", "5월 중순 #등산 스파이크 = 연휴 외생 이벤트"],
    )
    chart_slide(
        prs, next(pg), "이동평균 평활화", "05 · Smoothing", "02_moving_average.png",
        [f"포커스: #{FOCUS} (6월 모멘텀 1위)", "회색 = 원계열(노이즈)", "7일 MA = 주간 주기 상쇄", "28일 MA = 장기 추세만", "1~5월 꾸준한 우상향 확인"],
    )
    chart_slide(
        prs, next(pg), "정상성 진단과 차분", "05 · Stationarity", "03_stationarity.png",
        ["위: 롤링 평균이 계속 상승 → 평균 비불변 = 비정상", "롤링 표준편차도 5월에 점프", "아래: 1차 차분 후 평균이 0에 고정", "차분으로 (약)정상성 확보 → ARIMA의 d=1에 해당"],
    )
    chart_slide(
        prs, next(pg), "자기상관함수(ACF)", "05 · ACF", "04_acf.png",
        ["왼쪽(원계열): 천천히 감쇠 → 추세 지배 = 비정상 신호", "오른쪽(차분): lag 7·14·21 스파이크 → 주간 계절성", "점선 = ±1.96/√n 백색잡음 한계", "이론 그대로: 차분이 추세를 지우자 주기가 드러남"],
    )
    chart_slide(
        prs, next(pg), "가법 분해", "06 · Decomposition", "05_decomposition.png",
        ["X = 추세 + 계절 + 잔차로 분리", "추세: 1월~5월 +60% 수준 상승", "계절: 주말 양(+), 주중 음(−) 규칙 진동", "잔차: 5월 연휴 스파이크 외엔 백색잡음에 가까움"],
    )
    chart_slide(
        prs, next(pg), "주간 계절성 프로파일", "06 · Seasonality", "06_weekly_profile.png",
        ["요일 평균을 자기 평균=100으로 정규화", "#등산·#피크닉·#캠핑: 주말(토) 피크", "#다이어트: 월요일 피크 — '월요일 결심' 효과", "게시 타이밍 설계의 직접 근거"],
    )
    chart_slide(
        prs, next(pg), "성장 모멘텀", "06 · Momentum", "07_growth.png",
        [f"1월 첫 주=100 지수화", f"연초 대비: #피크닉 +{g['피크닉']:.0f}% · #캠핑 +{g['캠핑']:.0f}% · #등산 +{g['등산']:.0f}%",
         f"최근 30일: #등산 +{m['등산']:.1f}% · #클라이밍 +{m['클라이밍']:.1f}% 만 상승", "봄 태그(피크닉·캠핑)는 5월 피크 통과 후 하락 전환"],
    )
    chart_slide(
        prs, next(pg), "14일 예측", "06 · Forecast", "08_forecast.png",
        ["모형: 선형추세 + 요일 계절성 (최근 8주 적합)", "점선 = 6/11~6/24 예측 경로", "음영 = 95% 예측구간 (잔차 분산 기반)", "주말마다 솟는 요일 패턴까지 재현"],
    )

    # ---------------- 결론 ---------------- #
    section_slide(prs, "Conclusion", "그래서, 지금 뭘 써야 하나")

    cards_slide(
        prs, next(pg), "2026-06-11 기준 트렌딩 해시태그", "07 · Now trending",
        cards=[
            ("📈 지금 상승 중 (최근 30일 모멘텀 +)", "#등산  #클라이밍",
             f"등산 +{m['등산']:.1f}%, 클라이밍 +{m['클라이밍']:.1f}% — 초여름에도 모멘텀 유지, 지금 올라탈 태그", PINK),
            ("🌊 피크 통과 (계절성 하락 전환)", "#피크닉  #캠핑",
             f"연초 대비 +{g['피크닉']:.0f}%/+{g['캠핑']:.0f}%로 컸지만 최근 30일 {m['피크닉']:.0f}%/{m['캠핑']:.0f}% — 봄 수요 종료", PURPLE),
            ("💪 안정 수요 (시즌 진입)", "#다이어트",
             f"최근 30일 {m['다이어트']:+.1f}%로 하락 멈춤 — 여름 직전 수요 재점화 구간, 월요일 피크 활용", INK),
            ("📉 약세 지속", "#카페  #국내여행",
             f"연초 대비 {g['카페']:.0f}%/{g['국내여행']:.0f}% — 단독 태그보다 상승 태그와 조합 권장", GREY),
        ],
    )

    bullets_slide(
        prs, next(pg), "콘텐츠 활용 전략", "07 · Action",
        lead="시계열 인사이트 → 게시물 전략: 무엇을(태그), 언제(요일), 어떻게(조합)",
        bullets=[
            ("주력 태그", "#등산 #클라이밍 중심 콘텐츠 — 상승 모멘텀 구간에 게시해 노출 탄력 확보"),
            ("게시 타이밍", "아웃도어 태그는 금~토 오전(주말 피크 직전), #다이어트는 일 저녁~월 오전('월요일 결심' 수요 선점)"),
            ("태그 조합", "상승 태그 + 안정 태그 묶기: #등산 × #다이어트(여름 준비), #클라이밍 × #카페(실내 연계)"),
            ("리스크 관리", "#피크닉·#캠핑은 하락 전환 — 신규 기획 보류, 기존 콘텐츠는 내년 3~4월 재활용 예약"),
            ("운영 루프", "데이터랩 재수집(주 1회) → 모멘텀 갱신 → 태그 포트폴리오 리밸런싱"),
        ],
    )

    # ---------------- 부록 ---------------- #
    code_slide(
        prs, next(pg), "핵심 코드 ①  데이터 → 평활화 → 차분", "08 · Code",
        blocks=[
            ("tidy 패널을 날짜×태그 행렬로 (pandas pivot)",
             'wide = panel.pivot(index="date", columns="hashtag", values="ratio")'),
            ("이동평균 평활화 — 7일 중심 윈도우가 주간 주기를 상쇄",
             "ma7 = s.rolling(7, center=True).mean()   # 추세 + 노이즈 제거"),
            ("1차 차분으로 정상화 (ARIMA의 d=1)",
             "diff = s.diff()                          # X_t - X_{t-1}\n"
             "diff.rolling(28).mean()                  # 평균이 0에 붙는지 진단"),
            ("가법 분해 — 추세는 이동평균, 계절은 '같은 요일끼리 평균'",
             "trend    = s.rolling(7, center=True).mean()\n"
             'seasonal = (s - trend).groupby(s.index.weekday).transform("mean")\n'
             "residual = s - trend - seasonal          # 백색잡음이면 성공"),
        ],
    )

    code_slide(
        prs, next(pg), "핵심 코드 ②  ACF → 예측", "08 · Code",
        blocks=[
            ("표본 ACF 직접 구현 — ρ(k) = γ(k)/γ(0)",
             "x = x - x.mean()\n"
             "rho_k = np.dot(x[:n-k], x[k:]) / np.dot(x, x)\n"
             "band  = 1.96 / np.sqrt(n)                # 백색잡음 95% 한계"),
            ("선형추세 적합 (최소제곱) + 요일 더미 계절성",
             "coef = np.polyfit(t, s.values, 1)        # 기울기·절편\n"
             "seasonal = pd.Series(s.values - np.polyval(coef, t)) \\\n"
             "             .groupby(s.index.weekday).mean()"),
            ("14일 외삽 + 95% 예측구간 (잔차 표준편차 기반)",
             "future = np.polyval(coef, future_t) + seasonal[future_weekday]\n"
             "lo, hi = future - 1.96*resid_std, future + 1.96*resid_std"),
        ],
    )

    bullets_slide(
        prs, next(pg), "한계와 재현", "08 · Limitations",
        lead="정직한 캐비앗: 이 분석이 말할 수 있는 것과 없는 것",
        bullets=[
            ("프록시 한계", "네이버 검색량 ≠ 인스타그램 게시량 — '한국 사용자의 주제 관심도'의 대리지표로 해석해야 함"),
            ("상대값 한계", "데이터랩 ratio는 요청 내 정규화 — 태그 간 절대 규모 비교 불가, 시간 변화·모멘텀 비교만 유효"),
            ("실시간성", "2026-06-11 수집 스냅샷 — '실시간 스트림'이 아니라 주기 재수집으로 갱신하는 구조 (전일까지 일별 제공)"),
            ("모형 한계", "선형추세+요일 더미는 외생 이벤트(연휴 스파이크)·추세 전환을 늦게 반영 — SARIMA·지수평활이 다음 단계"),
            ("재현", "pip install -r presentation/requirements.txt && python presentation/build_deck.py — CSV 고정으로 결정론적 재생성"),
        ],
    )

    prs.save(OUT)
    return OUT, len(prs.slides._sldIdLst)


if __name__ == "__main__":
    path, n = build()
    print(f"Saved: {path} ({n} slides)")
